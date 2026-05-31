from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_chinese_font(run, font_name="宋体"):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_style(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=True):
    paragraph.alignment = alignment
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.3)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        set_chinese_font(run)
        run.font.size = Pt(18 if level == 0 else 16 if level == 1 else 14)
        run.font.bold = True
    return heading


def add_paragraph(document, text, first_line_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = document.add_paragraph(text)
    set_paragraph_style(paragraph, alignment=align, first_line_indent=first_line_indent)
    for run in paragraph.runs:
        set_chinese_font(run)
        run.font.size = Pt(12)
    return paragraph


def add_formula(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_chinese_font(run)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    return paragraph


def add_table(document, rows, columns, title=None):
    if title:
        add_paragraph(document, title, first_line_indent=False)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, col in enumerate(columns):
        header[i].text = col
    for row in rows:
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            value = row.get(col, "")
            if isinstance(value, float):
                value = f"{value:.4e}"
            cells[i].text = str(value)
    for row_idx, table_row in enumerate(table.rows):
        for cell in table_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_chinese_font(run)
                    run.font.size = Pt(10.5)
                    if row_idx == 0:
                        run.font.bold = True
    return table


def add_picture(document, path, caption, width=6.0):
    path = Path(path)
    if not path.exists():
        add_paragraph(document, f"[图片不存在: {path.name}]", first_line_indent=False)
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_para = document.add_paragraph(caption)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_para.runs:
        set_chinese_font(run)
        run.font.size = Pt(10.5)


def _metric_rows(summary_rows):
    rows = []
    for row in summary_rows:
        rows.append(
            {
                "模型": row["model"],
                "平均相对L2误差": row["relative_l2_mean"],
                "相对L2标准差": row["relative_l2_std"],
                "MSE": row["mse"],
                "u分量相对L2": row["u_relative_l2_mean"],
                "v分量相对L2": row["v_relative_l2_mean"],
                "参数量": row["parameters"],
            }
        )
    return rows


def _sample_rows(per_sample_rows):
    return [
        {"模型": row["model"], "测试样本": row["sample"], "相对L2误差": row["relative_l2"]}
        for row in per_sample_rows
    ]


def create_word_report(report_path, cfg, data_summary, summary_rows, per_sample_rows, assets_dir, output_dir):
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    assets_dir = Path(assets_dir)
    output_dir = Path(output_dir)

    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.styles["Normal"].font.size = Pt(12)

    add_heading(document, "科学计算第三次作业：Cavity Flow中的DeepONet与FNO对比", level=0)
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run("神经算子方法：DeepONet 与 Fourier Neural Operator")
    set_chinese_font(run)
    run.font.size = Pt(12)

    add_heading(document, "摘要", level=1)
    add_paragraph(
        document,
        "本报告围绕二维腔流（Cavity Flow）稳态速度场预测问题，分别实现了PyTorch版DeepONet与二维Fourier Neural Operator（FNO）。"
        "两个模型均以顶部壁面速度边界条件u_bc为输入，输出整个65×65网格上的水平速度u与垂直速度v。"
        "实验使用相同训练集与测试集，并通过整体相对L2误差、MSE、u/v分量误差和可视化结果进行公平比较。"
        "报告同时说明了仓库中已有DeepONet ODE demo和FNO earthquake 1D demo的作用与依赖。"
        "关键词：神经算子；DeepONet；Fourier Neural Operator；Cavity Flow；PyTorch。",
        first_line_indent=False,
    )

    add_heading(document, "1  作业目标与仓库背景", level=1)
    add_paragraph(
        document,
        "第三次作业关注函数到函数映射的学习问题。当前仓库中已有两个神经算子示例："
        "Codes/Neural_Operator/DeepONet提供TensorFlow v1风格的ODE算子demo，"
        "Codes/Neural_Operator/FNO_earthquake1D提供PyTorch一维FNO示例。"
        "本作业的主实验对象为Data/Cavity中的腔流数据，因此需要在同一数据集上实现DeepONet与FNO并进行对比。"
    )

    add_heading(document, "2  数据集说明", level=1)
    add_paragraph(
        document,
        "训练数据来自Cavity_Flow.mat，测试数据来自Cavity_Flow_Test.mat。输入变量u_bc表示上壁边界速度，"
        "输出变量u_data和v_data分别表示稳态速度场的两个分量，x_2d与y_2d给出空间网格坐标。"
        "由于数据文件未显式提供归一化参数，本文根据训练集统计量对输入和输出进行标准化，并在评估阶段反标准化回物理量。"
    )
    display_rows = []
    for row in data_summary:
        display_rows.append(
            {
                "文件": Path(row["file"]).name,
                "字段": row["key"],
                "shape": row["shape"],
                "dtype": row["dtype"],
                "最小值": row["min"],
                "最大值": row["max"],
            }
        )
    add_table(document, display_rows, ["文件", "字段", "shape", "dtype", "最小值", "最大值"], "表1  Cavity数据字段统计")

    add_heading(document, "3  方法原理", level=1)
    add_heading(document, "3.1  DeepONet", level=2)
    add_paragraph(
        document,
        "DeepONet将算子学习拆分为branch网络与trunk网络。branch网络读取离散边界函数u_bc，"
        "trunk网络读取查询坐标(x,y)，二者的隐变量逐项相乘并求和，得到该空间点处的速度分量。"
    )
    add_formula(document, "G_theta(u_bc)(x,y) = sum_k B_k(u_bc) T_k(x,y) + b")

    add_heading(document, "3.2  Fourier Neural Operator", level=2)
    add_paragraph(
        document,
        "FNO将输入场提升到高维通道空间，在频域中对低频模态进行线性变换，再通过反傅里叶变换回到物理空间。"
        "本实现将u_bc扩展到二维网格，并拼接x、y坐标通道，输出两个速度通道(u,v)。"
    )
    add_formula(document, "v_{l+1}(x) = sigma(W v_l(x) + F^{-1}(R_l · F(v_l))(x))")

    add_heading(document, "3.3  训练目标与评价指标", level=2)
    add_paragraph(document, "两个模型均以标准化后的速度场均方误差作为训练损失：")
    add_formula(document, "L = mean(||G_theta(u_bc) - [u, v]||_2^2)")
    add_paragraph(document, "测试阶段在反标准化后的物理量上计算相对L2误差与MSE：")
    add_formula(document, "Relative L2 = ||u_pred - u_true||_2 / ||u_true||_2")

    add_heading(document, "4  实验设置", level=1)
    setup_rows = [
        {"项目": "DeepONet", "设置": f"latent={cfg['deeponet']['latent_dim']}, width={cfg['deeponet']['branch_width']}, depth={cfg['deeponet']['depth']}"},
        {"项目": "FNO", "设置": f"modes=({cfg['fno']['modes1']},{cfg['fno']['modes2']}), width={cfg['fno']['width']}, depth={cfg['fno']['depth']}"},
        {"项目": "训练", "设置": f"batch={cfg['training']['batch_size']}, lr={cfg['training']['learning_rate']}, DeepONet epoch={cfg['training']['deeponet_epochs']}, FNO epoch={cfg['training']['fno_epochs']}"},
        {"项目": "设备", "设置": cfg.get("_device", "auto")},
    ]
    add_table(document, setup_rows, ["项目", "设置"], "表2  实验超参数")

    add_heading(document, "5  实验结果", level=1)
    add_table(
        document,
        _metric_rows(summary_rows),
        ["模型", "平均相对L2误差", "相对L2标准差", "MSE", "u分量相对L2", "v分量相对L2", "参数量"],
        "表3  DeepONet与FNO测试误差对比",
    )
    add_table(document, _sample_rows(per_sample_rows), ["模型", "测试样本", "相对L2误差"], "表4  逐测试样本相对L2误差")

    add_picture(document, assets_dir / "loss_curves.png", "图1  DeepONet与FNO训练/测试损失曲线")
    add_picture(document, assets_dir / "metric_comparison.png", "图2  DeepONet与FNO误差指标对比")
    add_picture(document, assets_dir / "deeponet_u_heatmap.png", "图3  DeepONet的u分量真值、预测与误差热力图")
    add_picture(document, assets_dir / "deeponet_v_heatmap.png", "图4  DeepONet的v分量真值、预测与误差热力图")
    add_picture(document, assets_dir / "deeponet_speed_heatmap.png", "图5  DeepONet速度模长真值、预测与误差")
    add_picture(document, assets_dir / "fno_u_heatmap.png", "图6  FNO的u分量真值、预测与误差热力图")
    add_picture(document, assets_dir / "fno_v_heatmap.png", "图7  FNO的v分量真值、预测与误差热力图")
    add_picture(document, assets_dir / "fno_speed_heatmap.png", "图8  FNO速度模长真值、预测与误差")
    add_picture(document, assets_dir / "deeponet_streamlines.png", "图9  DeepONet预测速度场流线图")
    add_picture(document, assets_dir / "fno_streamlines.png", "图10  FNO预测速度场流线图")

    add_heading(document, "6  结论", level=1)
    best = min(summary_rows, key=lambda r: r["relative_l2_mean"])
    add_paragraph(
        document,
        f"从本次训练结果看，{best['model']}取得了更低的平均相对L2误差。DeepONet的优势在于显式建模输入函数与查询点之间的关系，"
        "适合点查询形式的算子学习；FNO则利用频域卷积在规则网格上学习全局空间相关性，适合二维场预测。"
        "两种模型均能从上壁边界条件学习到腔流速度场的主要结构，误差热力图和速度模长图进一步展示了模型在局部区域的拟合差异。"
    )
    add_paragraph(document, f"所有数值结果、图像和模型文件均保存在：{output_dir}", first_line_indent=False)
    document.save(report_path)
    return report_path
